"""Create the standard job-list workbook used by this Codex Harness.

Run: python scripts/create_job_workbook.py [output.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def main() -> None:
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("jobs/待投岗位.xlsx")
    output.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    jobs = book.active
    jobs.title = "待投岗位"
    headers = ["URL", "模式", "启用", "优先级", "指定简历", "备注"]
    jobs.append(headers)
    jobs.append(["https://example.com/job/123", "定制", "是", 10, "", "请替换为真实岗位"])
    profile = book.create_sheet("个人资料")
    profile_headers = ["字段键", "问题别名", "答案", "类型", "敏感级别", "允许AI改写", "备注"]
    profile.append(profile_headers)
    profile.append(["self_intro", "自我介绍/个人优势", "", "开放题", "普通", "是", "由 Codex 草拟后确认"])
    for sheet in (jobs, profile):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
    book.save(output)
    print(output.resolve())


if __name__ == "__main__":
    main()

